import base64
import io
import logging
from langchain_core.messages import HumanMessage
from langchain_core.tools import tool
from psycopg.rows import dict_row

from app.db.db import get_conn
from app.db.storage import BUCKET, get_storage_client
from app.utils.model_provider import get_llm

logger = logging.getLogger(__name__)


def _extract_pdf(data: bytes) -> str:
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    parts: list[str] = []
    llm = None

    for page_num, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if text:
            parts.append(f"[Page {page_num}]\n{text}")

        # Pages with little text but images → describe images via LLM
        image_list = page.get_images(full=True)
        if image_list and len(text) < 100:
            if llm is None:
                llm = get_llm()
            for img_index, img_info in enumerate(image_list):
                xref = img_info[0]
                base_image = doc.extract_image(xref)
                img_bytes = base_image["image"]
                b64 = base64.b64encode(img_bytes).decode()
                ext = base_image["ext"]
                msg = HumanMessage(
                    content=[
                        {"type": "text", "text": "Describe this image concisely for document indexing."},
                        {"type": "image_url", "image_url": {"url": f"data:image/{ext};base64,{b64}"}},
                    ]
                )
                description = llm.invoke([msg]).content
                parts.append(f"[Page {page_num} Image {img_index + 1}]\n{description}")

    return "\n\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = " ".join(run.text for run in para.runs).strip()
                    if line:
                        texts.append(line)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"[Notes] {notes}")
        if texts:
            parts.append(f"[Slide {slide_num}]\n" + "\n".join(texts))

    return "\n\n".join(parts)


def _extract_image(data: bytes, mime: str = "jpeg") -> str:
    llm = get_llm()
    b64 = base64.b64encode(data).decode()
    msg = HumanMessage(
        content=[
            {"type": "text", "text": "Describe this image in detail for document indexing and retrieval."},
            {"type": "image_url", "image_url": {"url": f"data:image/{mime};base64,{b64}"}},
        ]
    )
    return llm.invoke([msg]).content


@tool(parse_docstring=True)
async def extract_text(document_id: str) -> str:
    """Download a document from storage and extract its full text content.

    Use this as the FIRST tool call after update_status("processing"). Call
    exactly once per document. The returned text must be passed directly to
    chunk_and_embed — do not modify or truncate it.

    Supports: pdf, pptx, md, txt, image. For PDF and PPTX, page/slide numbers
    are included as markers (e.g. [Page 3]). Image-heavy pages are described
    via LLM. For standalone image files the entire content is an LLM description.

    Args:
        document_id: UUID of the document to extract. Use exactly as provided.

    Returns:
        Full extracted text of the document, ready for chunking.

    Raises:
        RuntimeError: If document_id not found in database.
        ValueError: If file_type is not supported.
    """
    async with get_conn() as conn:
        async with conn.cursor(row_factory=dict_row) as cur:
            await cur.execute(
                "SELECT file_type, storage_path FROM documents WHERE id = %s",
                (document_id,),
            )
            row = await cur.fetchone()

    if not row:
        raise RuntimeError(f"Document {document_id} not found")

    file_type = row["file_type"].lower()
    storage_path = row["storage_path"]

    client = get_storage_client()
    data: bytes = client.storage.from_(BUCKET).download(storage_path)

    if file_type == "pdf":
        return _extract_pdf(data)
    elif file_type == "pptx":
        return _extract_pptx(data)
    elif file_type in ("md", "txt"):
        return data.decode("utf-8", errors="replace")
    elif file_type == "image":
        return _extract_image(data)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")
